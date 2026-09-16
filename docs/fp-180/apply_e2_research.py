# -*- coding: utf-8 -*-
"""FP-180 — First research pass: unblock the E2 multicloud scenario.

E2 (multicloud consolidation and price comparison) is the scenario where
comparing platforms is the central contribution of FP-51, and it held a single
scored pair (nOps). The coverage analysis showed the cheapest path to unblock
the rest: four platforms needed between one and four cells each to cross the
70% coverage threshold.

This pass fills those cells with evidence investigated against official
product documentation on 2026-09-16, quoted and cited per cell.

Scope rule applied when writing a value
---------------------------------------
Scenario-independent indicators (V2, O2, AU2, M1, M2, AD2, P1, D1, D2 —
properties of the product) are written to every scenario of the product, since
one investigation supports them everywhere. Scenario-dependent indicators
(V1, A1, A2, O1, AU1, I1, I2, AD1) are written only to E2, because the evidence
was read against E2's inputs and outputs.

Not every target cell could be filled. Finout D2 has no official documentation
of proprietary dependencies or an exit path, and Finout AU1 is not described
either way in the official docs, so both stay NE and Finout stays below the
threshold. That is the methodology working: no value is invented to close a gap.

Run after mark_scenario_inherited.py. Idempotent.

Usage: python apply_e2_research.py
"""
from __future__ import annotations

import collections
import datetime as dt
import shutil
import sys
from pathlib import Path

import openpyxl

from apply_scope_corrections import (BASE_W, SENS_W, evaluate, read_scores,
                                     rebuild_charts, write_results_sheet)

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE / "FP-180_matriz_comparativa.xlsx"
REPORT = HERE / "investigacion-e2.md"
VERIFIED = "2026-09-16"

SCENARIO_INDEPENDENT = {"V2", "O2", "AU2", "M1", "M2", "AD2", "P1", "D1", "D2"}

# (product, indicator) -> (value, evidence type, evidence text, source)
RESEARCH: dict[tuple[str, str], tuple[object, str, str, str]] = {
    ("CloudHealth", "AD1"): (
        2, "Oficial",
        "La plataforma define tres roles integrados y permite derivarlos: Administrator "
        "«has access to all privileges across all data»; Power Users «have the ability to "
        "perform all operations available to an administrator except the ability to create, "
        "edit, or delete organizations and users»; y Standard Users «can view but not modify "
        "most content». Además, «you can copy one of the built-in role documents (Standard "
        "User, Administrator, or Power User) and modify the copy to create your own. A role "
        "document is a set of privileges (read, write, update, delete)». Las organizaciones "
        "acotan qué datos ve cada usuario, lo que permite separar al responsable FinOps "
        "central, a arquitectura empresarial y a compras/finanzas exigidos por E2. No se "
        "asigna 3 porque falta la segunda evidencia (uso verificable) que pide FP-177.",
        "https://techdocs.broadcom.com/us/en/vmware-tanzu/cloudhealth/tanzu-cloudhealth/saas/"
        "tnz-cloudhealth/using-and-managing-tanzu-cloudhealth-managing-classic-organizations-"
        "if-applicable.html"),

    ("CloudZero", "V2"): (
        2, "Oficial",
        "La asignación reconcilia contra la factura: «The remaining spend goes to a default "
        "element so the total always reconciles with the actual bill». Para multinube, "
        "AnyCost y el Common Bill Format ingieren datos de cualquier proveedor, y el tipo de "
        "costo facturado permite la conciliación: «Billed Cost reflects the exact prices you "
        "will be invoiced for, and you can use it when you are reconciling against your "
        "actual invoice». Existen conexiones nativas a Azure (MCA, EA, CSP) y GCP (export de "
        "BigQuery), de modo que el total es rastreable hasta la fuente de facturación de cada "
        "proveedor.",
        "https://docs.cloudzero.com/docs/costformation-allocating-shared-costs ; "
        "https://docs.cloudzero.com/docs/anycost-common-bill-format-cbf ; "
        "https://docs.cloudzero.com/docs/cost-types"),

    ("CloudZero", "I2"): (
        2, "Oficial",
        "Integración documentada con el flujo operativo del escenario: Jira — «Create Work "
        "Items from Optimize Recommendations and Anomalies, then track their Jira status in "
        "CloudZero»; Slack — «Receive cost notifications directly in Slack channels, routed "
        "by View»; Google Chat; y Snowflake — «Share CloudZero cost data to your Snowflake "
        "account for SQL-based analysis», que sostiene el reporte consolidado que E2 espera. "
        "Es integración definida y reproducible, no una exportación manual.",
        "https://docs.cloudzero.com/docs/integrations ; "
        "https://docs.cloudzero.com/docs/notifications"),

    ("CloudZero", "AD2"): (
        2, "Oficial",
        "Mecanismo de ownership documentado: «Views combine a grouping, filters, and "
        "notification channels into a single preset, giving each team ownership of the costs "
        "they are responsible for», y cada View «monitors its scope for cost changes and "
        "sends updates and anomaly alerts to the configured Slack channel or email». La "
        "habilitación cuenta con material oficial de onboarding en CloudZero Academy. Proceso "
        "y guía definidos; no se asigna 3 porque no hay evidencia de adopción medida.",
        "https://docs.cloudzero.com/docs/cloudzero ; "
        "https://academy.cloudzero.com/docs/onboarding-1"),

    ("Finout", "V2"): (
        2, "Oficial",
        "MegaBill permite descender del total al detalle: «You can quickly drill down into "
        "the costs/usage by clicking an area on the graph», y la Cost API V2 expone consulta "
        "directa de costo y uso con filtros y group-by, de modo que el detalle es exportable y "
        "reproducible. Límite registrado para E2: la documentación oficial declara que «Cost "
        "data in MegaBill is available for all cloud providers and services supported by "
        "Finout. Usage data is available for AWS, GCP, and Azure», es decir, el dato de uso no "
        "cubre a todos los proveedores. No se encontró documentación oficial de conciliación "
        "contra la factura del proveedor; la afirmación de conciliación patentada aparece solo "
        "en material comercial y por eso no se usa para puntuar.",
        "https://docs.finout.io/user-guide/inform/megabill ; "
        "https://docs.finout.io/api/finout-api/finout-api-v2"),

    ("Finout", "AD2"): (
        2, "Oficial",
        "Gobernanza documentada como función del producto: la gobernanza de etiquetas "
        "«allows you to define tagging policies across multiple cloud environments, offering "
        "visibility into non-compliant resources and costs while simplifying resource "
        "tracking», que es un mecanismo de seguimiento para uso sostenido en un entorno "
        "multinube. La documentación incluye además la guía de incorporación de nuevos "
        "usuarios a la cuenta. Proceso definido; sin evidencia de adopción medida para 3.",
        "https://docs.finout.io/user-guide/operate/tag-governance ; "
        "https://docs.finout.io/get-started-with-finout/introduction-to-finouts-suite-of-features"),

    ("Vantage", "V2"): (
        2, "Oficial",
        "El informe de costos permite descender del total al detalle: «Within the table below "
        "the graph, you can drill down into your report to create precise filtering and "
        "grouping views». El Data Dictionary documenta los campos del dato de costo y VQL "
        "(Vantage Query Language) permite consultas reproducibles sobre ellos. La procedencia "
        "queda marcada en el propio registro: para proveedores soportados «Vantage tags every "
        "row with an [...] billing_source value [...] so you can filter, group, and allocate "
        "costs by data source».",
        "https://docs.vantage.sh/cost_reports ; https://docs.vantage.sh/data_dictionary ; "
        "https://docs.vantage.sh/vql"),

    ("Vantage", "A2"): (
        2, "Oficial",
        "Regla de reparto de costo compartido documentada y configurable: «With percent-based "
        "cost allocation, you can filter Cost Reports to show back shared resources, like "
        "support costs or multi-tenant databases, to the team or department that uses them». "
        "Las asignaciones encadenadas están documentadas: una Virtual Tag basada en costo "
        "«can now aggregate by the key of another allocated Virtual Tag in the same provider "
        "allocation chain», lo que permite repartos de varios pasos reproducibles.",
        "https://docs.vantage.sh/vantage_university_cost_allocation ; "
        "https://docs.vantage.sh/virtual_tagging/"),

    ("Vantage", "M2"): (
        2, "Oficial",
        "Las diferencias semánticas entre proveedores se hacen explícitas y se normalizan: "
        "«Vantage normalizes Usage Unit labels in the console for consistent spelling and "
        "capitalization. For example, Vantage standardized Hours, hrs, Hrs, and hours to "
        "Hours», y se aplica un esquema de etiquetas normalizado por proveedor para costos de "
        "IA. La cobertura se declara de forma explícita —la lista de proveedores soportados "
        "para reporte por uso está enumerada en la documentación— y los no soportados se "
        "canalizan por Custom Providers: «you can upload costs from providers that don't "
        "expose billing APIs, aren't yet supported by Vantage, or originate from custom "
        "systems». El límite de cobertura queda visible, que es lo que mide M2.",
        "https://docs.vantage.sh/usage_based_reporting ; "
        "https://docs.vantage.sh/connecting_custom_providers ; "
        "https://docs.vantage.sh/data_dictionary"),

    ("Vantage", "AD2"): (
        2, "Oficial",
        "Habilitación documentada por rol en Vantage University, con rutas para FinOps "
        "Analyst, Developers y Finance Managers, y una ruta específica de asignación de "
        "costos. El ownership se sostiene en segmentos jerárquicos: los segmentos de "
        "asignación «let you create hierarchical structures for your costs, mapping them to "
        "various organizational entities, such as business units, teams, or services», con "
        "jerarquías anidadas por equipo y subequipo. Proceso y guía definidos; sin evidencia "
        "de adopción medida para 3.",
        "https://docs.vantage.sh/vantage_university ; "
        "https://docs.vantage.sh/vantage_university_cost_allocation ; "
        "https://docs.vantage.sh/segments"),
}

# Target cells that could not be filled, with the reason. Recorded, not scored.
NOT_FILLED = {
    ("Finout", "D2"): "No se encontró documentación oficial que describa dependencias "
                      "propietarias ni un camino de salida o configuración reversible. La "
                      "documentación de API cubre el acceso programático al dato (D1), no la "
                      "reversibilidad del acoplamiento. Se mantiene NE.",
    ("Finout", "AU1"): "La documentación oficial de CostGuard describe generación de "
                       "recomendaciones y permisos de gestión, pero no describe ejecución ni "
                       "bloqueo de acciones de costo. Como la ausencia no está afirmada "
                       "explícitamente, corresponde NE y no 0. Ruta alternativa evaluada para "
                       "desbloquear Finout; tampoco se pudo cerrar con evidencia disponible.",
}

NOTE = ("Investigación E2 FP-180 ({date}): celda completada con evidencia propia del escenario, "
        "verificada contra documentación oficial.")


def apply_research(sheet) -> list[dict]:
    header = [c.value for c in sheet[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    changes = []
    for row in range(2, sheet.max_row + 1):
        scenario = sheet.cell(row, col["Escenario"]).value
        product = sheet.cell(row, col["Producto"]).value
        indicator = sheet.cell(row, col["Indicador"]).value
        target = RESEARCH.get((product, indicator))
        if not target or scenario == "VALIDACION":
            continue
        if indicator not in SCENARIO_INDEPENDENT and scenario != "E2":
            continue                      # scenario-dependent evidence belongs to E2 only

        cell = sheet.cell(row, col["Valor (0-3/NE/NA)"])
        if str(cell.value).strip().upper() != "NE":
            continue                      # only fills gaps; never overwrites a scored cell

        value, evidence_type, evidence, source = target
        cell.value = value
        sheet.cell(row, col["Tipo de evidencia"]).value = evidence_type
        sheet.cell(row, col["Evidencia"]).value = evidence
        sheet.cell(row, col["Fuente"]).value = source
        sheet.cell(row, col["Fecha de verificacion"]).value = VERIFIED
        note_cell = sheet.cell(row, col["Notas"])
        note = NOTE.format(date=VERIFIED)
        note_cell.value = f"{note_cell.value} | {note}" if note_cell.value else note
        changes.append({"esc": scenario, "prod": product, "ind": indicator, "val": value})
    return changes


def write_report(changes, before, after) -> None:
    lines = [
        "# FP-180 — Primera pasada de investigación: escenario E2",
        "",
        f"> Generado por `apply_e2_research.py` el {VERIFIED}. Evidencia verificada contra "
        "documentación oficial de cada producto. Queda a validación humana.",
        "",
        "## Por qué E2 primero",
        "",
        "E2 —consolidación multinube y comparación de precios— es el escenario donde comparar "
        "plataformas es el aporte central de FP-51, y tenía **un solo par puntuado** (nOps). "
        "El análisis de cobertura mostró que era también el más barato de desbloquear: cuatro "
        "plataformas necesitaban entre una y cuatro celdas para cruzar el umbral del 70%.",
        "",
        "## Regla de alcance aplicada al escribir cada valor",
        "",
        "Los indicadores independientes del escenario (V2, O2, AU2, M1, M2, AD2, P1, D1, D2 — "
        "propiedades del producto) se escriben en todos los escenarios del producto, porque una "
        "sola investigación los sostiene. Los dependientes del escenario (V1, A1, A2, O1, AU1, "
        "I1, I2, AD1) se escriben solo en E2, porque la evidencia se leyó contra las entradas y "
        "salidas de E2.",
        "",
        f"## Celdas completadas ({len({(c['prod'], c['ind']) for c in changes})} únicas, "
        f"{len(changes)} filas)",
        "",
        "| Producto | Indicador | Valor | Fuente |",
        "|---|---|---:|---|",
    ]
    seen = set()
    for c in changes:
        key = (c["prod"], c["ind"])
        if key in seen:
            continue
        seen.add(key)
        source = RESEARCH[key][3].split(" ; ")[0]
        lines.append(f"| {c['prod']} | {c['ind']} | {c['val']} | {source} |")

    lines += ["", "## Celdas que no se pudieron completar", "",
              "| Producto | Indicador | Motivo |", "|---|---|---|"]
    for (prod, ind), reason in NOT_FILLED.items():
        lines.append(f"| {prod} | {ind} | {reason} |")

    lines += ["", "## Efecto en los resultados", "",
              "| Escenario | Producto | Antes | Después | Cobertura antes | Cobertura después |",
              "|---|---|---|---|---:|---:|"]
    for key in before:
        old, new = before[key], after[key]
        if old["score"] == new["score"] and abs(old["coverage"] - new["coverage"]) < 0.01:
            continue
        fmt = lambda r: (f"{r['score']:.2f} ({r['band']})" if r["score"] is not None
                         else "Insufficient evidence")
        lines.append(f"| {key[0]} | {key[2]} | {fmt(old)} | {fmt(new)} | "
                     f"{old['coverage']:.1f}% | {new['coverage']:.1f}% |")

    e2_after = sum(1 for k, r in after.items() if k[0] == "E2" and r["score"] is not None)
    lines += [
        "",
        "## Nota sobre Finout",
        "",
        "Finout recibió dos de las tres celdas que necesitaba (V2 y AD2), pero **sigue por "
        "debajo del umbral**. Su tercer criterio más barato era Dependencia, y no existe "
        "documentación oficial de dependencias propietarias ni de camino de salida; la ruta "
        "alternativa (Automatización, vía CostGuard) tampoco se pudo cerrar, porque la "
        "documentación no describe ejecución ni bloqueo de acciones. Ninguna de las dos se "
        "rellenó con un valor inventado.",
        "",
        f"Tras esta pasada, E2 tiene {e2_after} pares puntuados.",
        "",
    ]
    REPORT.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not WORKBOOK.exists():
        print(f"ERROR: no se encontro {WORKBOOK}")
        return 1

    shutil.copy2(WORKBOOK, WORKBOOK.with_suffix(".xlsx.bak"))
    workbook = openpyxl.load_workbook(WORKBOOK)
    scores = workbook["Puntuacion"]

    before, _ = read_scores(scores)
    before_results = {k: evaluate(v, BASE_W) for k, v in before.items()}

    changes = apply_research(scores)
    print(f"Celdas completadas: {len(changes)} filas")

    after, _ = read_scores(scores)
    after_results = {k: evaluate(v, BASE_W) for k, v in after.items()}

    write_results_sheet(workbook, after)
    rebuild_charts(workbook)
    workbook.save(WORKBOOK)
    if changes:
        write_report(changes, before_results, after_results)
        print(f"Informe: {REPORT.name}")
    else:
        print(f"Sin celdas pendientes; se conserva {REPORT.name} tal como estaba.")

    computable = sum(1 for r in after_results.values() if r["score"] is not None)
    e2 = sum(1 for k, r in after_results.items() if k[0] == "E2" and r["score"] is not None)
    print(f"Pares con puntaje calculable: {computable} de {len(after_results)} (E2: {e2})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
