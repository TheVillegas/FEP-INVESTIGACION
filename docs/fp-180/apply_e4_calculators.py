# -*- coding: utf-8 -*-
"""FP-180 — Close E4 with the two pricing calculators.

E4 (estimation and policy before deployment) had two scored pairs, Infracost and
Cloud Custodian, while the two alternatives registered by FP-126 — the AWS and
Azure pricing calculators — stayed below the coverage threshold.

This pass fills what was blocking them, from official documentation consulted on
2026-09-16, and fixes an inconsistency of my own: Optimizacion had been left NE
for both calculators when the absence is in fact verified. Neither tool detects
waste or rightsizing opportunities; they estimate cost for a configuration you
describe. That is a 0 under the FP-177 anchor, not a gap in the research, and
O2 follows as NA because there is no recommendation of its own to track.

AWS Pricing Calculator turns out to document its coupling and exit path, which
is rare in this population: it states where estimates are stored, for how long
the shared links stay valid, and offers delete operations through its API.

Usage: python apply_e4_calculators.py
"""
from __future__ import annotations

import datetime as dt
import shutil
import sys
from pathlib import Path

import openpyxl

from apply_scope_corrections import (BASE_W, evaluate, read_scores,
                                     rebuild_charts, write_results_sheet)

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE / "FP-180_matriz_comparativa.xlsx"
REPORT = HERE / "investigacion-e4.md"
VERIFIED = "2026-09-16"
SCENARIO = "E4"

RESEARCH = {
    ("AWS Pricing Calculator", "D2"): (
        2, "Oficial",
        "El acoplamiento y la salida estan documentados de forma inusualmente explicita para "
        "esta poblacion. La dependencia se declara: «Estimates are saved to the AWS public "
        "servers», y antes de compartir se exige leer un «Public server acknowledgment». El "
        "limite temporal tambien: «Estimate links created on or after May 31, 2023, remain "
        "valid for one year». Y la reversibilidad esta cubierta por operaciones de borrado "
        "documentadas en la API (DeleteWorkloadEstimate y BatchDeleteWorkloadEstimateUsage).",
        "https://docs.aws.amazon.com/pricing-calculator/latest/userguide/save-share-estimate.html ; "
        "https://docs.aws.amazon.com/aws-cost-management/latest/APIReference/"
        "API_AWSBCMPricingCalculator_DeleteWorkloadEstimate.html"),

    ("AWS Pricing Calculator", "O1"): (
        0, "Oficial (ausencia verificada)",
        "La herramienta estima el costo de una configuracion que el usuario modela; la "
        "documentacion no describe deteccion de desperdicio, rightsizing ni oportunidades de "
        "compromiso. Es una ausencia verificada frente al ancla del indicador, no una falta de "
        "investigacion: esa funcion corresponde a otros productos de la poblacion.",
        "https://docs.aws.amazon.com/pricing-calculator/latest/userguide/what-is-pricing-calculator.html"),

    ("AWS Pricing Calculator", "O2"): (
        "NA", "Justificacion de alcance",
        "Al no generar recomendaciones de optimizacion propias (ver O1), no aplica evaluar la "
        "explicacion de impacto ni el seguimiento de una recomendacion inexistente. Mismo "
        "criterio aplicado a AU2 cuando AU1 es 0.", ""),

    ("Azure Pricing Calculator", "D1"): (
        2, "Oficial",
        "La estimacion es exportable en formatos utilizables y reproducibles: la calculadora "
        "permite exportar las estimaciones a Excel, y las paginas de detalle de precios "
        "permiten exportar y guardar como CSV. Las vistas de Cost Management se descargan en "
        "PNG, Excel y CSV.",
        "https://learn.microsoft.com/en-us/azure/cost-management-billing/costs/pricing-calculator ; "
        "https://learn.microsoft.com/en-us/azure/cost-management-billing/costs/save-share-views"),

    ("Azure Pricing Calculator", "I2"): (
        1, "Oficial",
        "La integracion con el flujo de revision es por guardado, enlace compartido y "
        "exportacion: la calculadora permite guardar estimaciones y compartirlas —con inicio de "
        "sesion— y exportarlas a Excel o CSV para distribuirlas. No se documenta integracion "
        "con CI/CD ni con un flujo de aprobacion previo al despliegue, por lo que corresponde "
        "el nivel manual. Mismo criterio aplicado a AWS Pricing Calculator.",
        "https://learn.microsoft.com/en-us/azure/cost-management-billing/costs/pricing-calculator"),

    ("Azure Pricing Calculator", "O1"): (
        0, "Oficial (ausencia verificada)",
        "La calculadora convierte uso previsto en costo estimado; la documentacion no describe "
        "deteccion de desperdicio ni recomendaciones de optimizacion. Ausencia verificada "
        "frente al ancla, no falta de investigacion.",
        "https://learn.microsoft.com/en-us/azure/cost-management-billing/costs/pricing-calculator"),

    ("Azure Pricing Calculator", "O2"): (
        "NA", "Justificacion de alcance",
        "Al no generar recomendaciones de optimizacion propias (ver O1), no aplica evaluar su "
        "explicacion ni su seguimiento.", ""),
}


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not WORKBOOK.exists():
        print(f"ERROR: no se encontro {WORKBOOK}")
        return 1

    shutil.copy2(WORKBOOK, WORKBOOK.with_suffix(".xlsx.bak"))
    workbook = openpyxl.load_workbook(WORKBOOK)
    sheet = workbook["Puntuacion"]
    header = [c.value for c in sheet[1]]
    col = {h: i + 1 for i, h in enumerate(header)}

    changes = []
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, col["Escenario"]).value != SCENARIO:
            continue
        product = sheet.cell(row, col["Producto"]).value
        indicator = sheet.cell(row, col["Indicador"]).value
        target = RESEARCH.get((product, indicator))
        if not target:
            continue
        cell = sheet.cell(row, col["Valor (0-3/NE/NA)"])
        if str(cell.value).strip().upper() != "NE":
            continue
        value, ev_type, evidence, source = target
        cell.value = value
        sheet.cell(row, col["Tipo de evidencia"]).value = ev_type
        sheet.cell(row, col["Evidencia"]).value = evidence
        sheet.cell(row, col["Fuente"]).value = source
        sheet.cell(row, col["Fecha de verificacion"]).value = VERIFIED if source else ""
        note_cell = sheet.cell(row, col["Notas"])
        note = (f"Investigacion propia de E4 ({VERIFIED}): celda verificada contra las entradas "
                "y salidas del escenario de estimacion previa al despliegue.")
        note_cell.value = f"{note_cell.value} | {note}" if note_cell.value else note
        changes.append({"prod": product, "ind": indicator, "val": value})

    print(f"Celdas completadas en E4: {len(changes)}")

    after, _ = read_scores(sheet)
    after_results = {k: evaluate(v, BASE_W) for k, v in after.items()}
    write_results_sheet(workbook, after)
    rebuild_charts(workbook)
    workbook.save(WORKBOOK)

    lines = [
        "# FP-180 — Escenario E4, estimacion y politica antes del despliegue",
        "",
        f"> Generado por `apply_e4_calculators.py` el {VERIFIED}. Queda a validacion humana.",
        "",
        "## Por que",
        "",
        "E4 tenia dos pares puntuados —Infracost y Cloud Custodian— mientras las dos "
        "alternativas que FP-126 registro para la categoria, las calculadoras de precios de AWS "
        "y Azure, quedaban bajo el umbral de cobertura.",
        "",
        "## Una inconsistencia propia corregida",
        "",
        "El criterio Optimizacion habia quedado en `NE` para ambas calculadoras cuando la "
        "ausencia esta en realidad verificada: ninguna detecta desperdicio ni oportunidades de "
        "rightsizing, porque estiman el costo de una configuracion que el usuario describe. "
        "Eso es un `0` frente al ancla de FP-177, no un hueco de investigacion, y O2 pasa a "
        "`NA` porque no hay recomendacion propia que explicar ni seguir.",
        "",
        "## Un hallazgo sobre AWS Pricing Calculator",
        "",
        "Es de los pocos productos de la poblacion que documenta su acoplamiento y su salida. "
        "Declara donde se guardan las estimaciones («Estimates are saved to the AWS public "
        "servers»), exige un reconocimiento explicito antes de compartirlas, fija la vigencia "
        "de los enlaces en un ano y ofrece operaciones de borrado por API. Junto con Kubecost, "
        "son los unicos dos casos de `D2` puntuado en toda la matriz.",
        "",
        f"## Celdas completadas ({len(changes)})",
        "",
        "| Producto | Indicador | Valor |",
        "|---|---|---:|",
    ]
    for c in sorted(changes, key=lambda x: (x["prod"], x["ind"])):
        lines.append(f"| {c['prod']} | {c['ind']} | {c['val']} |")

    lines += ["", "## Estado de E4 tras esta pasada", "",
              "| Producto | S linea base | Banda | Cobertura |", "|---|---:|---|---:|"]
    scored = [(k, v) for k, v in after_results.items()
              if k[0] == SCENARIO and v["score"] is not None]
    for k, r in sorted(scored, key=lambda kv: -kv[1]["score"]):
        lines.append(f"| {k[2]} | {r['score']:.2f} | {r['band']} | {r['coverage']:.1f}% |")

    unscored = sorted((k[2], v["coverage"]) for k, v in after_results.items()
                      if k[0] == SCENARIO and v["score"] is None)
    if unscored:
        lines += ["", "Siguen sin puntaje en E4: " +
                  ", ".join(f"{n} ({c:.1f}%)" for n, c in unscored) + "."]
    lines.append("")
    REPORT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Informe: {REPORT.name}")

    e4 = [(k, v) for k, v in after_results.items()
          if k[0] == SCENARIO and v["score"] is not None]
    print(f"\nE4 con puntaje ({len(e4)}):")
    for k, r in sorted(e4, key=lambda kv: -kv[1]["score"]):
        print(f"   {k[2]:<30} {r['score']:>6.2f}  {r['band']:<8} cob {r['coverage']:.1f}%")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
