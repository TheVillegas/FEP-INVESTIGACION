# -*- coding: utf-8 -*-
"""FP-180 — Score E6, large distributed and scientific workloads.

E6 had no scored pair. Its critical criteria are Visibilidad and Integracion,
and its expected outputs include the components of total cost and the cost of
network per workflow — network being what sets this scenario apart, since the
ATLAS/CERN evidence behind it shows transfer costs can dominate.

Strict rule applied here: V1 and I1 are the two indicators where E6 differs from
every other scenario, and they are also its critical criteria. They are scored
only with evidence of their own about network and transfer cost. A product
without that evidence keeps them at NE, even if its general breakdown capability
is documented elsewhere — precisely so the critical criteria of this scenario
are never satisfied by a capability that was verified for another one.

The remaining scenario-dependent indicators are subsumed from the product's
origin scenario, with the argument recorded per cell.

Usage: python apply_e6_distributed.py
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl

from apply_scope_corrections import (BASE_W, evaluate, read_scores,
                                     rebuild_charts, write_results_sheet)

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE / "FP-180_matriz_comparativa.xlsx"
REPORT = HERE / "investigacion-e6.md"
VERIFIED = "2026-09-16"
SCENARIO = "E6"

ORIGIN_BY_CATEGORY = {"Cloud nativa": "E1", "Multicloud": "E2"}
SUBSUMED = ["A1", "A2", "O1", "AU1", "I2", "AD1", "P2"]   # V1 and I1 need their own evidence

SUBSUMPTION_NOTE = (
    "Evidencia subsumida desde {origin} ({date}): la capacidad no cambia por la escala ni por "
    "la distribucion de la carga. Lo que si distingue a E6 —el costo de red y transferencia por "
    "flujo de trabajo— se evalua en V1 e I1, que solo se puntuan con evidencia propia de este "
    "escenario. La fuente y la cita originales se conservan sin cambios.")

E6_OWN_EVIDENCE = {
    ("Vantage", "V1"): (
        2, "Oficial",
        "Los componentes del costo de red quedan desglosados por flujo: «For any "
        "network-enabled resource, Network Flow Reports provide visibility by source and "
        "destination to the flows within your network that are driving costs», y «each flow "
        "shows the estimated cost associated with that specific traffic route». Ademas "
        "reasigna la bolsa agregada de AWS a su origen real: los cargos «initially categorized "
        "under transit or egress in EC2-Other» se asignan «to the individual resource that "
        "created the charge».",
        "https://docs.vantage.sh/network_flow_reports ; https://docs.vantage.sh/cost_reports"),

    ("Vantage", "I1"): (
        2, "Oficial",
        "Integra los datos de red y transferencia que este escenario exige como entrada, "
        "incluyendo el trafico de cluster: «Transmitted bytes are grouped into mutually "
        "exclusive buckets. Each bucket is priced by the destination and, for S3-bound "
        "traffic, by how the source subnet egresses», lo que permite atribuir el costo de "
        "transferencia por destino.",
        "https://docs.vantage.sh/kubernetes_network_costs ; "
        "https://docs.vantage.sh/network_flow_reports"),

    ("AWS Cost Explorer", "V1"): (
        2, "Oficial",
        "El CUR expone los componentes de transferencia de datos de forma granular y "
        "documentada: los cargos se identifican por la columna lineItem/UsageType, con tipos "
        "distintos para transferencia entre regiones "
        "(Source Region-Destination Region-AWS-Out-Bytes), entre zonas de disponibilidad de una "
        "misma region (Region-DataTransfer-Regional-Bytes) y para CloudFront, distinguible por "
        "lineItem/ProductCode. Eso permite separar el componente de red del costo total, que es "
        "la salida que E6 espera.",
        "https://docs.aws.amazon.com/cur/latest/userguide/cur-data-transfers-charges.html"),

    ("AWS Cost Explorer", "I1"): (
        2, "Oficial",
        "La ingesta de los datos que pide el escenario —computo, almacenamiento y red— esta "
        "documentada en el propio CUR, que entrega una fila por linea de facturacion con sus "
        "columnas de producto y de item, incluidos los tipos de uso de transferencia de datos.",
        "https://docs.aws.amazon.com/cur/latest/userguide/Lineitem-columns.html ; "
        "https://docs.aws.amazon.com/cur/latest/userguide/cur-data-transfers-charges.html"),
}


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not WORKBOOK.exists():
        print(f"ERROR: no se encontro {WORKBOOK}")
        return 1

    shutil.copy2(WORKBOOK, WORKBOOK.with_suffix(".xlsx.bak"))
    workbook = openpyxl.load_workbook(WORKBOOK)
    sheet = workbook["Puntuacion"]
    header = [c.value for c in sheet[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    idx = {h: i for i, h in enumerate(header)}

    source = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        source[(row[idx["Producto"]], row[idx["Escenario"]], row[idx["Indicador"]])] = (
            row[idx["Valor (0-3/NE/NA)"]], row[idx["Tipo de evidencia"]],
            row[idx["Evidencia"]], row[idx["Fuente"]])

    changes = []
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, col["Escenario"]).value != SCENARIO:
            continue
        category = sheet.cell(row, col["Categoria"]).value
        product = sheet.cell(row, col["Producto"]).value
        indicator = sheet.cell(row, col["Indicador"]).value
        cell = sheet.cell(row, col["Valor (0-3/NE/NA)"])
        if str(cell.value).strip().upper() != "NE":
            continue

        own = E6_OWN_EVIDENCE.get((product, indicator))
        if own:
            value, ev_type, evidence, src = own
            note = (f"Investigacion propia de E6 ({VERIFIED}): celda verificada contra el "
                    "componente que distingue a este escenario, el costo de red y "
                    "transferencia por flujo de trabajo.")
        elif indicator in SUBSUMED:
            origin = ORIGIN_BY_CATEGORY.get(category)
            if not origin:
                continue
            got = source.get((product, origin, indicator))
            if not got or str(got[0]).strip().upper() in ("NE", "NA", "NONE"):
                continue
            value, ev_type, evidence, src = got
            note = SUBSUMPTION_NOTE.format(origin=origin, date=VERIFIED)
        else:
            continue

        cell.value = value
        sheet.cell(row, col["Tipo de evidencia"]).value = ev_type
        sheet.cell(row, col["Evidencia"]).value = evidence
        sheet.cell(row, col["Fuente"]).value = src
        sheet.cell(row, col["Fecha de verificacion"]).value = VERIFIED if src else ""
        note_cell = sheet.cell(row, col["Notas"])
        note_cell.value = f"{note_cell.value} | {note}" if note_cell.value else note
        changes.append({"prod": product, "ind": indicator, "val": value})

    print(f"Celdas completadas en E6: {len(changes)}")

    after, _ = read_scores(sheet)
    after_results = {k: evaluate(v, BASE_W) for k, v in after.items()}
    write_results_sheet(workbook, after)
    rebuild_charts(workbook)
    workbook.save(WORKBOOK)

    lines = [
        "# FP-180 — Escenario E6, carga distribuida y cientifica de gran escala",
        "",
        f"> Generado por `apply_e6_distributed.py` el {VERIFIED}. Queda a validacion humana.",
        "",
        "## Por que",
        "",
        "E6 no tenia ningun par puntuado. Sus criterios criticos son Visibilidad e Integracion, "
        "y entre sus salidas esperadas estan los componentes del costo total y el costo de red "
        "por flujo de trabajo. La red es lo que distingue a este escenario: la evidencia "
        "ATLAS/CERN que lo sustenta muestra que la transferencia puede dominar el costo.",
        "",
        "## Regla estricta aplicada",
        "",
        "V1 e I1 son los dos indicadores donde E6 se diferencia de cualquier otro escenario, y "
        "son ademas sus criterios criticos. Aqui se puntuan **solo con evidencia propia sobre "
        "costo de red y transferencia**. Un producto sin esa evidencia los conserva en `NE`, "
        "aunque su capacidad general de desglose este documentada en otro lado — precisamente "
        "para que los criterios criticos de este escenario nunca queden satisfechos por una "
        "capacidad verificada para otro.",
        "",
        "El resto de indicadores dependientes se subsume desde el escenario de origen del "
        "producto, con el argumento registrado celda por celda.",
        "",
        f"## Celdas completadas ({len(changes)})",
        "",
        "| Producto | Indicador | Valor | Tipo |",
        "|---|---|---:|---|",
    ]
    for c in sorted(changes, key=lambda x: (x["prod"], x["ind"])):
        kind = ("propia de E6" if (c["prod"], c["ind"]) in E6_OWN_EVIDENCE else "subsumida")
        lines.append(f"| {c['prod']} | {c['ind']} | {c['val']} | {kind} |")

    lines += ["", "## Estado de E6 tras esta pasada", "",
              "| Producto | Categoria | S linea base | Banda | Cobertura |",
              "|---|---|---:|---|---:|"]
    scored = [(k, v) for k, v in after_results.items()
              if k[0] == SCENARIO and v["score"] is not None]
    for k, r in sorted(scored, key=lambda kv: -kv[1]["score"]):
        lines.append(f"| {k[2]} | {k[1]} | {r['score']:.2f} | {r['band']} | "
                     f"{r['coverage']:.1f}% |")

    unscored = sorted((k[2], v["coverage"]) for k, v in after_results.items()
                      if k[0] == SCENARIO and v["score"] is None)
    if unscored:
        lines += ["", "## Siguen sin puntaje en E6", "",
                  "Les falta evidencia propia sobre costo de red y transferencia, que este "
                  "escenario exige en sus dos criterios criticos.", "",
                  "| Producto | Cobertura |", "|---|---:|"]
        for name, cov in unscored:
            lines.append(f"| {name} | {cov:.1f}% |")
    lines.append("")
    REPORT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Informe: {REPORT.name}")

    e6 = [(k, v) for k, v in after_results.items()
          if k[0] == SCENARIO and v["score"] is not None]
    print(f"\nE6 con puntaje ({len(e6)}):")
    for k, r in sorted(e6, key=lambda kv: -kv[1]["score"]):
        print(f"   {k[2]:<30} {r['score']:>6.2f}  {r['band']:<8} cob {r['coverage']:.1f}%")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
