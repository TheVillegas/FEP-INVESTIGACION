# -*- coding: utf-8 -*-
"""FP-180 — Mark inherited scenario cells as NE.

Problem this fixes
------------------
Every product was investigated once and its scores were copied into each of its
eligible scenarios: the 15 multi-scenario products carry identical values *and
identical evidence text* across all of them. The 61 rows of the calculation
sheet therefore hold 17 evaluations, not 61.

That contradicts FP-177 section 1, which sets the unit of result at product x
scenario and requires testing each product "con las mismas entradas del
escenario". A recommendation per scenario cannot come out of a matrix whose
numbers do not change per scenario.

What this pass does
-------------------
An indicator is scenario-dependent when its observable question in the FP-177
operational matrix refers to the scenario itself:

    V1  breakdown unit (account/project vs namespace vs tenant)
    A1  allocation entity of the scenario
    A2  shared-cost split rules (central to E5, marginal to E1)
    O1  waste "relevante al escenario" (literal wording)
    AU1 the cost actions of the scenario
    I1  ingests "los datos requeridos ... del escenario" (literal wording)
    I2  the "flujo de trabajo relevante" of the scenario (literal wording)
    AD1 access "para los roles del escenario" (literal wording)
    P2  contextual estimation/comparison of the scenario

The other nine (V2, O2, AU2, M1, M2, AD2, P1, D1, D2) describe properties of
the product itself, not of the scenario, so a single investigation supports
them across every scenario and they are left untouched.

For the scenario-dependent ones, the recorded evidence supports the value only
in the scenario where it was actually gathered — the anchor scenario of the
product's category under FP-177 section 4. Everywhere else the value is
inherited, has no evidence of its own, and becomes NE: it drops out of the
score and reduces coverage, which is what FP-177 prescribes for something not
evidenced. It does not become NA, which would renormalise the weights and hide
the gap.

Expect coverage to fall sharply. That is the point: it makes the real size of
the pending investigation visible instead of resolving it with copied values.

Run after apply_scope_corrections.py. Idempotent.

Usage: python mark_scenario_inherited.py
"""
from __future__ import annotations

import collections
import datetime as dt
import shutil
import sys
from pathlib import Path

import openpyxl

from apply_scope_corrections import (BASE_W, evaluate, read_scores,
                                     rebuild_charts, write_results_sheet)

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE / "FP-180_matriz_comparativa.xlsx"
REPORT = HERE / "celdas-heredadas.md"
TODAY = dt.date.today().isoformat()

SCENARIO_DEPENDENT = ["V1", "A1", "A2", "O1", "AU1", "I1", "I2", "AD1", "P2"]

# FP-177 section 4 assigns each category the scenario where it is the main actor;
# that is where the product's evidence was gathered.
ANCHOR_BY_CATEGORY = {
    "Cloud nativa": "E1",
    "Multicloud": "E2",
    "Kubernetes": "E3",
    "Estimacion temprana": "E4",
}

NOTE = ("Celda heredada FP-180 ({date}): el valor {old} proviene de la investigacion del "
        "escenario {anchor} y no cuenta con evidencia propia de {scenario}. Segun FP-177 el "
        "indicador {indicator} depende del escenario, por lo que pasa a NE (no evidenciado, "
        "reduce cobertura) hasta investigarlo. La evidencia y la fuente originales se conservan "
        "como referencia del escenario de origen.")


def mark_inherited(sheet) -> list[dict]:
    header = [c.value for c in sheet[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    changes = []
    for row in range(2, sheet.max_row + 1):
        scenario = sheet.cell(row, col["Escenario"]).value
        category = sheet.cell(row, col["Categoria"]).value
        indicator = sheet.cell(row, col["Indicador"]).value
        if scenario == "VALIDACION" or indicator not in SCENARIO_DEPENDENT:
            continue
        anchor = ANCHOR_BY_CATEGORY.get(category)
        if not anchor or scenario == anchor:
            continue                                   # evidence belongs to this scenario

        cell = sheet.cell(row, col["Valor (0-3/NE/NA)"])
        current = str(cell.value).strip().upper()
        if current in ("NE", "NA"):
            continue                                   # nothing inherited to strip

        sheet.cell(row, col["Tipo de evidencia"]).value = "Pendiente"
        note_cell = sheet.cell(row, col["Notas"])
        note = NOTE.format(date=TODAY, old=cell.value, anchor=anchor, scenario=scenario,
                           indicator=indicator)
        note_cell.value = f"{note_cell.value} | {note}" if note_cell.value else note
        changes.append({"escenario": scenario, "categoria": category,
                        "producto": sheet.cell(row, col["Producto"]).value,
                        "indicador": indicator, "anterior": cell.value, "ancla": anchor})
        cell.value = "NE"
    return changes


def write_report(changes, before, after) -> None:
    by_product = collections.defaultdict(list)
    for change in changes:
        by_product[change["producto"]].append(change)

    lines = [
        "# FP-180 — Celdas heredadas marcadas como NE",
        "",
        f"> Generado por `mark_scenario_inherited.py` el {TODAY}. "
        "Propuesta de corrección metodológica pendiente de validación del equipo.",
        "",
        "## Por qué",
        "",
        "Cada producto se investigó una sola vez y su puntuación se copió a todos sus escenarios "
        "elegibles: los 15 productos multiescenario tenían valores **y texto de evidencia "
        "idénticos** en todos ellos. Las 61 filas de la hoja Calculo contenían 17 evaluaciones, "
        "no 61.",
        "",
        "FP-177 §1 fija la unidad de resultado en producto × escenario y exige probar cada producto "
        "«con las mismas entradas del escenario». Una recomendación por escenario no puede salir de "
        "una matriz cuyos números no cambian entre escenarios.",
        "",
        "## Qué indicadores dependen del escenario",
        "",
        "Un indicador depende del escenario cuando su pregunta observable en la matriz de FP-177 "
        "se refiere al escenario mismo:",
        "",
        "| Indicador | Por qué depende del escenario |",
        "|---|---|",
        "| V1 | La unidad de desglose cambia: cuenta/proyecto en E1, namespace en E3, entidad en E5. |",
        "| A1 | La entidad de asignación es la del escenario. |",
        "| A2 | Las reglas de reparto compartido son centrales en E5 y marginales en E1. |",
        "| O1 | Pregunta por desperdicio «relevante al escenario» (texto literal). |",
        "| AU1 | Las acciones de costo son las del escenario. |",
        "| I1 | Pregunta si integra «los datos requeridos … del escenario» (texto literal). |",
        "| I2 | Pregunta por el «flujo de trabajo relevante» del escenario (texto literal). |",
        "| AD1 | Pregunta por acceso «para los roles del escenario» (texto literal). |",
        "| P2 | La estimación o comparación es contextual al escenario. |",
        "",
        "Los otros nueve (V2, O2, AU2, M1, M2, AD2, P1, D1, D2) describen propiedades del producto "
        "y no del escenario, así que una sola investigación los sostiene en todos y **no se "
        "tocaron**.",
        "",
        "## Escenario ancla por categoría",
        "",
        "La evidencia registrada sostiene el valor solo en el escenario donde se recogió: el "
        "escenario en que FP-177 §4 sitúa a esa categoría como actor principal.",
        "",
        "| Categoría | Escenario ancla |",
        "|---|---|",
    ]
    for category, anchor in ANCHOR_BY_CATEGORY.items():
        lines.append(f"| {category} | {anchor} |")

    lines += [
        "",
        "Fuera del ancla el valor es heredado, no tiene evidencia propia y pasa a `NE`: sale del "
        "puntaje y reduce la cobertura, que es lo que FP-177 prescribe para lo no evidenciado. No "
        "pasa a `NA`, que renormalizaría los pesos y escondería el hueco.",
        "",
        f"## Celdas marcadas ({len(changes)})",
        "",
        "| Producto | Escenario | Indicador | Valor heredado | Escenario de origen |",
        "|---|---|---|---:|---|",
    ]
    for product in sorted(by_product):
        for change in sorted(by_product[product],
                             key=lambda c: (c["escenario"], c["indicador"])):
            lines.append(f"| {product} | {change['escenario']} | {change['indicador']} | "
                         f"{change['anterior']} | {change['ancla']} |")

    lines += ["", "## Efecto en los resultados", "",
              "| Escenario | Producto | S base antes | S base después | Cobertura antes | Cobertura después |",
              "|---|---|---:|---:|---:|---:|"]
    for key in before:
        old, new = before[key], after[key]
        if old["score"] == new["score"] and abs(old["coverage"] - new["coverage"]) < 0.01:
            continue
        fmt = lambda r: (f"{r['score']:.2f}" if r["score"] is not None else "Insuf.")
        lines.append(f"| {key[0]} | {key[2]} | {fmt(old)} | {fmt(new)} | "
                     f"{old['coverage']:.1f}% | {new['coverage']:.1f}% |")

    computable = sum(1 for k in after if after[k]["score"] is not None)
    lines += [
        "",
        "## Lectura de este resultado",
        "",
        f"Quedan {computable} de {len(after)} pares con puntaje calculable. No es una pérdida de "
        "trabajo: es la medida real de lo investigado. La evidencia recogida sostiene una "
        "evaluación por producto en su escenario de origen, y esta pasada deja de presentar como "
        "61 evaluaciones lo que son 17.",
        "",
        f"El trabajo pendiente queda cuantificado: {len(changes)} celdas por investigar con "
        "evidencia propia del escenario. Cada una necesita documentación oficial que responda la "
        "pregunta observable del indicador **para las entradas y salidas de ese escenario**, no "
        "para el producto en general.",
        "",
    ]
    REPORT.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if not WORKBOOK.exists():
        print(f"ERROR: no se encontro {WORKBOOK}")
        return 1

    backup = WORKBOOK.with_suffix(".xlsx.bak")
    shutil.copy2(WORKBOOK, backup)

    workbook = openpyxl.load_workbook(WORKBOOK)
    scores = workbook["Puntuacion"]

    before, _ = read_scores(scores)
    before_results = {k: evaluate(v, BASE_W) for k, v in before.items()}

    changes = mark_inherited(scores)
    print(f"Celdas heredadas marcadas como NE: {len(changes)}")

    after, _ = read_scores(scores)
    after_results = {k: evaluate(v, BASE_W) for k, v in after.items()}

    write_results_sheet(workbook, after)
    rebuild_charts(workbook)
    workbook.save(WORKBOOK)
    if changes:
        write_report(changes, before_results, after_results)
        print(f"Informe: {REPORT.name}")
    else:
        # Nothing left to mark: keep the report of the pass that did the work.
        print(f"Sin celdas heredadas pendientes; se conserva {REPORT.name} tal como estaba.")

    computable = sum(1 for r in after_results.values() if r["score"] is not None)
    print(f"Pares con puntaje calculable: {computable} de {len(after_results)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
